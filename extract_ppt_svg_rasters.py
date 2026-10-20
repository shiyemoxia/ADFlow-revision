"""Recover raster images embedded inside SVG shape fills in a PPTX slide.

Some Figure 6 source pictures are stored as data-URI images inside SVG fills,
so a search restricted to ordinary PowerPoint picture shapes misses them.
This script only extracts evidence; it does not edit the presentation.
"""

import argparse
import base64
import csv
import hashlib
from io import BytesIO
from pathlib import Path
from xml.etree import ElementTree

from PIL import Image
from pptx import Presentation


REL_EMBED = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed'
XLINK_HREF = '{http://www.w3.org/1999/xlink}href'
SVG_IMAGE = '{http://www.w3.org/2000/svg}image'
EXTENSIONS = {'image/png': '.png', 'image/jpeg': '.jpg'}


def extract(pptx_path, slide_number, output_dir):
    presentation = Presentation(pptx_path)
    if not 1 <= slide_number <= len(presentation.slides):
        raise ValueError('Slide number is outside the presentation')
    slide = presentation.slides[slide_number - 1]
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    rows = []
    saved = {}

    for shape_number, shape in enumerate(slide.shapes, start=1):
        for node in shape._element.iter():
            if not node.tag.endswith('svgBlip'):
                continue
            relation_id = node.get(REL_EMBED)
            if not relation_id:
                continue
            svg_part = slide.part.related_part(relation_id)
            root = ElementTree.fromstring(svg_part.blob)
            for embedded_index, raster in enumerate(root.iter(SVG_IMAGE), start=1):
                uri = raster.get(XLINK_HREF, '')
                if not uri.startswith('data:') or ';base64,' not in uri:
                    continue
                header, encoded = uri.split(',', 1)
                mime_type = header[5:].split(';', 1)[0]
                if mime_type not in EXTENSIONS:
                    continue
                payload = base64.b64decode(encoded, validate=True)
                key = (relation_id, embedded_index)
                if key not in saved:
                    filename = f'slide{slide_number:02d}_{relation_id}_{embedded_index:02d}{EXTENSIONS[mime_type]}'
                    path = output_dir / filename
                    path.write_bytes(payload)
                    saved[key] = path
                with Image.open(BytesIO(payload)) as image:
                    width, height = image.size
                rows.append({
                    'slide': slide_number,
                    'shape': shape_number,
                    'shape_left_inches': f'{shape.left / 914400:.3f}',
                    'shape_top_inches': f'{shape.top / 914400:.3f}',
                    'svg_part': str(svg_part.partname),
                    'relation_id': relation_id,
                    'embedded_index': embedded_index,
                    'raster_file': saved[key].name,
                    'width': width,
                    'height': height,
                    'sha256': hashlib.sha256(payload).hexdigest(),
                })

    index_path = output_dir / f'slide{slide_number:02d}_svg_rasters.csv'
    with index_path.open('w', newline='', encoding='utf-8') as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]) if rows else [
            'slide', 'shape', 'shape_left_inches', 'shape_top_inches',
            'svg_part', 'relation_id', 'embedded_index', 'raster_file',
            'width', 'height', 'sha256',
        ])
        writer.writeheader()
        writer.writerows(rows)
    return len(saved), len(rows), index_path


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--pptx', default='实验图源.pptx')
    parser.add_argument('--slide', type=int, default=6)
    parser.add_argument('--output-dir', default='results_figure_matching/svg_rasters')
    args = parser.parse_args()
    unique, references, index_path = extract(args.pptx, args.slide, args.output_dir)
    print(f'Extracted {unique} unique raster images from {references} shape references')
    print(f'Index: {index_path}')


if __name__ == '__main__':
    main()
