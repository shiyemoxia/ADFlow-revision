"""Extract PPT reference pictures and rank possible original dataset images.

The ranking is an aid to manual verification, not proof of sample identity.
Requires: pip install python-pptx
"""

import argparse
import csv
from itertools import chain
from pathlib import Path

import numpy as np
from PIL import Image, ImageOps
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


def _fingerprint(image):
    image = ImageOps.fit(image.convert('RGB'), (48, 48), method=Image.Resampling.LANCZOS)
    return np.asarray(image, dtype=np.float32).reshape(-1) / 255.0


def _candidate_paths(mvtec_root, visa_root, loco_root):
    candidates = []
    for dataset, root in (('mvtec', mvtec_root), ('visa', visa_root),
                          ('mvtec-loco', loco_root)):
        if not root or not Path(root).is_dir():
            continue
        root = Path(root)
        if dataset == 'visa':
            images = chain(root.glob('*/Data/Images/Anomaly/*'),
                           root.glob('*/Data/Images/Normal/*'))
        else:
            images = root.glob('*/test/*/*')
        for path in images:
            if path.is_file() and path.suffix.lower() in ('.png', '.jpg', '.jpeg'):
                candidates.append((dataset, root, path))
    if not candidates:
        raise FileNotFoundError('No dataset test images found at the supplied roots')
    return candidates


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--pptx', default='实验图源.pptx')
    parser.add_argument('--mvtec-root', default='../mvtec-ad')
    parser.add_argument('--visa-root', default='../visa')
    parser.add_argument('--loco-root', default='../mvtec-loco')
    parser.add_argument('--slides', nargs='+', type=int, default=[4, 7],
                        help='one-based slide numbers (default: Figure 6 and Figure 9 source slides)')
    parser.add_argument('--output-dir', default='results_figure_matching')
    parser.add_argument('--top-k', type=int, default=5)
    parser.add_argument('--extract-only', action='store_true',
                        help='extract PPT pictures and positions before datasets are available')
    args = parser.parse_args()
    out = Path(args.output_dir)
    out.mkdir(parents=True, exist_ok=True)
    valid, vectors = [], None
    if not args.extract_only:
        candidates = _candidate_paths(args.mvtec_root, args.visa_root, args.loco_root)
        fingerprints = []
        for item in candidates:
            try:
                with Image.open(item[2]) as picture:
                    fingerprints.append(_fingerprint(picture))
                valid.append(item)
            except (OSError, ValueError):
                continue
        if not fingerprints:
            raise ValueError('No readable candidate images')
        vectors = np.stack(fingerprints)

    deck = Presentation(args.pptx)
    rows, references = [], []
    for slide_number in args.slides:
        if not 1 <= slide_number <= len(deck.slides):
            raise ValueError('Slide {} is outside this presentation'.format(slide_number))
        for shape_number, shape in enumerate(deck.slides[slide_number - 1].shapes, start=1):
            if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
                continue
            try:
                from io import BytesIO
                with Image.open(BytesIO(shape.image.blob)) as reference:
                    extension = reference.format.lower() if reference.format else 'png'
                    reference_vector = _fingerprint(reference)
                    reference_path = out / 'slide{:02d}_shape{:03d}.{}'.format(
                        slide_number, shape_number, extension)
                    reference_path.write_bytes(shape.image.blob)
            except (OSError, ValueError):
                continue
            references.append({
                'slide': slide_number, 'shape': shape_number,
                'reference_image': reference_path.name,
                'reference_left': shape.left, 'reference_top': shape.top,
                'reference_width': shape.width, 'reference_height': shape.height,
            })
            if args.extract_only:
                continue
            error = np.mean((vectors - reference_vector) ** 2, axis=1)
            for rank, candidate_index in enumerate(np.argsort(error)[:args.top_k], start=1):
                dataset, root, candidate = valid[int(candidate_index)]
                rows.append({
                    'slide': slide_number, 'shape': shape_number,
                    'reference_image': reference_path.name,
                    'reference_left': shape.left, 'reference_top': shape.top,
                    'rank': rank, 'dataset': dataset,
                    'dataset_relative_path': candidate.relative_to(root).as_posix(),
                    'distance': '{:.8f}'.format(error[candidate_index]),
                    'warning': 'Candidate only: visually verify input and GT before selecting a Figure 9 sample',
                })
    reference_csv = out / 'reference_images.csv'
    with reference_csv.open('w', newline='', encoding='utf-8') as handle:
        writer = csv.DictWriter(handle, fieldnames=[
            'slide', 'shape', 'reference_image', 'reference_left', 'reference_top',
            'reference_width', 'reference_height'])
        writer.writeheader()
        writer.writerows(references)
    print('Extracted {} PPT pictures to {}'.format(len(references), out))
    if args.extract_only:
        return
    output_csv = out / 'candidate_matches.csv'
    with output_csv.open('w', newline='', encoding='utf-8') as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]) if rows else [
            'slide', 'shape', 'reference_image', 'reference_left', 'reference_top',
            'rank', 'dataset', 'dataset_relative_path', 'distance', 'warning'])
        writer.writeheader()
        writer.writerows(rows)
    print('Saved {} candidate matches to {}'.format(len(rows), output_csv))
    print('Match by picture position and visual appearance; do not treat rankings as confirmed IDs.')


if __name__ == '__main__':
    main()
